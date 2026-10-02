"""Create one slide with a floorplan PNG and editable category legend."""

from __future__ import annotations

import os
from pathlib import Path
import subprocess

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt
except ImportError as exc:
    raise RuntimeError(
        "PowerPoint export requires python-pptx. Use .venv/bin/python main.py "
        "or install requirements.txt in your Python environment."
    ) from exc


EXPORT_UTIL = Path(__file__).resolve().parents[2] / "hw/syn/xilinx/xrt/export_util.tcl"
LEGEND_TCL = """namespace eval ::vortex_util {variable library_only 1}
if {[catch {
    source $::env(VORTEX_PHOTO_LEGEND_TCL)
    set specs [::vortex_util::category_specs]
    foreach key {simt memory mxu dma misc} {
        set spec [dict get $specs $key]
        puts "$key\t[dict get $spec label]\t[join [dict get $spec rgb] ,]"
    }
} message]} {
    puts stderr $message
    exit 1
}
"""


def load_legend() -> list[tuple[str, str, tuple[int, ...]]]:
    """Read the same Tcl definitions used to color the captured implementation."""
    result = subprocess.run(
        ["tclsh"], input=LEGEND_TCL, text=True, capture_output=True, check=True,
        env={**os.environ, "VORTEX_PHOTO_LEGEND_TCL": str(EXPORT_UTIL)}, timeout=15,
    )
    legend = []
    for line in result.stdout.splitlines():
        key, label, colors = line.split("\t")
        rgb = tuple(int(value) for value in colors.split(","))
        if len(rgb) != 3 or any(value < 0 or value > 255 for value in rgb):
            raise ValueError(f"Invalid legend RGB for {key}: {rgb}")
        legend.append((key, label, rgb))
    if len(legend) != 5:
        raise ValueError("Expected five Vortex categories in export_util.tcl")
    return legend


def add_text(slide, name: str, text: str, x: float, y: float, w: float, h: float,
             *, size: int = 20, bold: bool = False, color=(35, 45, 55)):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(text.splitlines()):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor(*color)
    return shape


def export_floorplan(image: Path, output: Path, title: str = "FPGA floorplan") -> Path:
    if not image.is_file():
        raise FileNotFoundError(f"Floorplan PNG does not exist: {image}")
    if image.suffix.lower() != ".png":
        raise ValueError("--image must be a PNG file")
    if output.suffix.lower() != ".pptx":
        raise ValueError("PowerPoint output must have a .pptx suffix")
    if output.exists():
        raise FileExistsError(f"PowerPoint output already exists: {output}")
    legend = load_legend()
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    add_text(slide, "Title", title, 0.65, 0.3, 12, 0.5, size=28, bold=True)

    # Fit inside the left column without cropping or stretching the PNG.
    picture = slide.shapes.add_picture(str(image), 0, 0, height=Inches(6.05))
    if picture.width > Inches(5.6):
        picture.height = int(picture.height * Inches(5.6) / picture.width)
        picture.width = Inches(5.6)
    picture.left = Inches(0.65) + (Inches(5.6) - picture.width) // 2
    picture.top = Inches(1.05) + (Inches(6.05) - picture.height) // 2
    picture.name = "Floorplan PNG"

    add_text(slide, "Legend heading", "Vortex categories", 6.75, 1.35, 5.8, 0.4,
             size=22, bold=True)
    for index, (key, label, rgb) in enumerate(legend):
        y = 2.0 + index * 0.72
        square = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(6.75), Inches(y + 0.07), Inches(0.30), Inches(0.30),
        )
        square.name = f"Legend swatch: {key}"
        square.fill.solid()
        square.fill.fore_color.rgb = RGBColor(*rgb)
        square.line.fill.background()
        add_text(slide, f"Legend label: {key}", label, 7.25, y, 5.45, 0.48, size=19)
    add_text(
        slide, "Overlay note",
        "Vivado resource/pblock colors are separate from this legend.\n"
        "Yellow HBM labels do not indicate MXU.",
        6.75, 6.10, 5.75, 0.65, size=12, color=(95, 105, 115),
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    print(f"Saved PowerPoint: {output}", flush=True)
    return output
